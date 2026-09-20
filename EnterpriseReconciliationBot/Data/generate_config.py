import openpyxl
from pathlib import Path

wb = openpyxl.Workbook()

# Sheet 1: Settings
ws_settings = wb.active
ws_settings.title = "Settings"
ws_settings.append(["Name", "Value", "Description"])
ws_settings.append(["SapBaseUrl", "http://localhost:8000/sap/opu/odata/sap/API_OP_INVOICE_PROCESS_SRV", "Base endpoint for SAP S/4HANA OData service"])
ws_settings.append(["InvoicesInputFolder", "Data/Input", "Folder containing vendor invoice files"])
ws_settings.append(["InvoicesProcessedFolder", "Data/Processed", "Folder for successfully posted invoices"])
ws_settings.append(["InvoicesExceptionsFolder", "Data/Exceptions", "Folder for rejected or manual review invoices"])
ws_settings.append(["MaxRetryNumber", "3", "Maximum number of retries for system exceptions"])
ws_settings.append(["ConfidenceThreshold", "0.85", "Minimum regex confidence threshold before AI fallback"])

# Sheet 2: Constants
ws_constants = wb.create_sheet(title="Constants")
ws_constants.append(["Name", "Value", "Description"])
ws_constants.append(["LogPrefix", "[EnterpriseReconciliationBot]", "Standard enterprise logging prefix"])
ws_constants.append(["VatRateStandard", "0.075", "Standard Nigerian VAT rate (7.5%)"])
ws_constants.append(["WhtRateServices", "0.05", "Standard WHT withholding on technical services (5%)"])
ws_constants.append(["VatToleranceThreshold", "0.05", "Maximum allowable rounding tolerance"])

# Sheet 3: Assets
ws_assets = wb.create_sheet(title="Assets")
ws_assets.append(["Name", "Asset", "Description"])
ws_assets.append(["SapCredentialAsset", "SAP_S4HANA_CREDENTIAL", "Orchestrator credential asset for SAP login"])
ws_assets.append(["GeminiApiKeyAsset", "Gemini_API_Key", "Orchestrator text asset for Google Gemini AI Fallback API"])
ws_assets.append(["NotificationEmailAsset", "RPA_Alerts_Email", "Orchestrator asset for finance exception alerts"])

target_dir = Path(__file__).parent
wb.save(target_dir / "Config.xlsx")
print("Config.xlsx created successfully at", target_dir / "Config.xlsx")
